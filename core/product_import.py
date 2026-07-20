# Import en masse de produits depuis un fichier Excel (.xlsx), et génération
# du modèle Excel à télécharger pour respecter le bon format de colonnes.

import io

import openpyxl
from openpyxl.styles import Font, PatternFill

# Colonnes attendues dans le fichier Excel, dans cet ordre. La ligne d'en-tête
# (première ligne) est toujours ignorée lors de l'import.
COLUMNS = ["Désignation", "Description", "Prix (F CFA)", "Stock", "Seuil critique"]


def build_import_template() -> io.BytesIO:
    """Construit un classeur Excel vierge avec les bons en-têtes et une
    ligne d'exemple, à proposer au téléchargement pour l'admin."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Produits"

    header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_index, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_index, value=header)
        cell.fill = header_fill
        cell.font = header_font

    # Ligne d'exemple, à remplacer par les vraies données
    ws.append(["Disque de frein ventilé (unité)", "Compatible véhicules courants", 15000, 20, 5])

    for col_index in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_index)].width = 26

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def parse_products_file(uploaded_file, company):
    """Lit le fichier Excel uploadé et retourne (produits_à_créer, erreurs).
    `produits_à_créer` est une liste d'instances Product non encore
    enregistrées (à créer via bulk_create par l'appelant). N'écrit rien en
    base ici : la validation et l'écriture sont séparées pour rester sûres."""
    from .models import Product  # import différé pour éviter tout cycle

    wb = openpyxl.load_workbook(uploaded_file, data_only=True)
    ws = wb.active

    products_to_create = []
    errors = []

    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name, description, price, stock, critical_threshold = (list(row) + [None] * 5)[:5]

        if name is None or str(name).strip() == "":
            continue  # ligne vide : ignorée silencieusement

        name = str(name).strip()

        try:
            price = float(price)
            if price <= 0:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(f"Ligne {row_index} : prix invalide pour « {name} ».")
            continue

        try:
            stock = int(stock) if stock is not None and str(stock).strip() != "" else 0
            if stock < 0:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(f"Ligne {row_index} : stock invalide pour « {name} ».")
            continue

        try:
            critical_threshold = (
                int(critical_threshold)
                if critical_threshold is not None and str(critical_threshold).strip() != ""
                else 5
            )
        except (TypeError, ValueError):
            critical_threshold = 5

        products_to_create.append(
            Product(
                company=company,
                name=name,
                description=(str(description).strip() if description else None),
                price=price,
                stock=stock,
                critical_threshold=critical_threshold,
            )
        )

    return products_to_create, errors
